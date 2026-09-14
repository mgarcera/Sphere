#!/usr/bin/env python3
"""
The same frames, as a .pptx Canva can import — so the layout arrives as EDITABLE ELEMENTS
rather than as a flat PNG.

build.py is the source of truth. This file reads its LAYOUT, FRAMES and edge sampling and
emits one slide per frame carrying three separate objects:

    background   a full-bleed rectangle with the sampled top -> bottom gradient
    caption      a real text box, Newsreader, at the real size and leading
    device       the capture, pre-rounded and pre-shadowed as a transparent PNG

Canva has no connector this session, so the handoff is a file: canva/sphere-frames.pptx,
then canva.com -> Create a design -> Import file.

Two things do not survive the trip and are deliberate losses:
  - text-wrap: balance. Canva breaks lines greedily; re-break by hand if a caption goes ragged.
  - the shadow is baked into the device PNG's alpha. Moving the device moves its shadow, which
    is right; changing the shadow means re-running this script, not editing in Canva.

    python3 AppStore/frames/canva.py
"""
import pathlib, sys
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from lxml import etree

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from build import LAYOUT as L, FRAMES, CAP, W, H, edges

HERE = pathlib.Path(__file__).parent
OUT = HERE / "canva"
PX = 9525                                  # 1 px at 96 DPI, in EMU
SHADOW = dict(dy=40, blur=90, alpha=0.22)  # matches build.py's box-shadow
BLEED = 220                                # transparent room around the device for the shadow

def device_png(src, dest):
    """Capture -> rounded, shadowed, transparent PNG at exactly device_width across."""
    im = Image.open(src).convert("RGBA")
    w = L["device_width"]
    h = round(im.height * w / im.width)
    im = im.resize((w, h), Image.LANCZOS)

    mask = Image.new("L", (w, h), 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, w - 1, h - 1], L["device_radius"], fill=255)
    im.putalpha(mask)

    canvas = Image.new("RGBA", (w + BLEED * 2, h + BLEED * 2), (0, 0, 0, 0))
    shade = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    shade.paste((0, 0, 0, round(255 * SHADOW["alpha"])),
                (BLEED, BLEED + SHADOW["dy"]), mask)
    canvas.alpha_composite(shade.filter(ImageFilter.GaussianBlur(SHADOW["blur"] / 2)))
    canvas.alpha_composite(im, (BLEED, BLEED))
    canvas.save(dest)
    return canvas.size

def gradient(shape, top, bot):
    f = shape.fill
    f.gradient()
    s = f.gradient_stops
    s[0].color.rgb, s[0].position = RGBColor.from_string(top[1:]), 0.34
    s[1].color.rgb, s[1].position = RGBColor.from_string(bot[1:]), 1.0
    # DrawingML angles run clockwise from east in 60000ths of a degree; 90 deg = downwards.
    lin = shape.fill._xPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}gradFill") \
                        .find("{http://schemas.openxmlformats.org/drawingml/2006/main}lin")
    lin.set("ang", "5400000")
    shape.line.fill.background()

def caption(slide, text, ink):
    box = slide.shapes.add_textbox(Emu(L["pad"] * PX), Emu(L["caption_top"] * PX),
                                   Emu((W - 2 * L["pad"]) * PX), Emu(700 * PX))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.line_spacing = L["caption_leading"]
    r = p.add_run(); r.text = text
    r.font.name, r.font.size = "Newsreader", Pt(L["caption_size"] * 0.75)
    r.font.color.rgb = RGBColor.from_string(ink[1:])
    # letter-spacing: -.015em, in hundredths of a point
    r.font._rPr.set("spc", str(round(-0.015 * L["caption_size"] * 0.75 * 100)))

def main():
    OUT.mkdir(exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    blank = prs.slide_layouts[6]

    made, waiting = [], []
    for name, text, scheme in FRAMES:
        src = CAP / f"{name}.png"
        if not src.exists():
            waiting.append(name); continue
        top, bot = edges(Image.open(src))
        dw, dh = device_png(src, OUT / f"{name}-device.png")

        slide = prs.slides.add_slide(blank)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(W * PX), Emu(H * PX))
        gradient(bg, top, bot)
        caption(slide, text, "#F4F2EE" if scheme == "dark" else "#141317")
        slide.shapes.add_picture(str(OUT / f"{name}-device.png"),
                                 Emu(round((W / 2 - dw / 2) * PX)),
                                 Emu((L["device_top"] - BLEED) * PX),
                                 Emu(dw * PX), Emu(dh * PX))
        made.append((name, dw - BLEED * 2, dh - BLEED * 2))

    deck = OUT / "sphere-frames.pptx"
    prs.save(deck)
    print(f"  {deck}  {W}x{H}px  ({W/96:.3f} x {H/96:.3f} in)")
    for name, w, h in made:
        print(f"  {name:10} device {w}x{h}")
    if waiting:
        print(f"\n  waiting on a capture: {', '.join(waiting)}")

if __name__ == "__main__":
    main()
